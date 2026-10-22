#!/usr/bin/env python3
"""从任意 .pptx 剥出空白底版，供 build_pptx.py --template 使用。

用法:
    python make_base.py 你的PPT.pptx -o assets/base.pptx

原理: 删除演示文稿中的全部幻灯片，但保留母版、版式、主题与字体定义。
生成的课件因此继承这套视觉语言（配色、字体、占位符样式），
不需要重新猜一套配色。

注意: 幻灯片上的手工排版（直接放在页面上的文本框与图片）会随幻灯片一起被删除。
如果原稿的关键样式是"画在页面上"而不是定义在母版/版式里，剥出的底版不会带上 ——
这种情况下先打开原稿，把样式整理进母版，再跑本脚本。
"""

import argparse
import os
import sys

from pptx import Presentation


def strip_slides(src, out):
    prs = Presentation(src)
    n = len(prs.slides)

    xml_slides = prs.slides._sldIdLst
    for sld in list(xml_slides):
        prs.part.drop_rel(sld.rId)
        xml_slides.remove(sld)

    os.makedirs(os.path.dirname(os.path.abspath(out)) or ".", exist_ok=True)
    prs.save(out)

    return n, len(prs.slide_layouts), len(prs.slide_masters)


def main():
    ap = argparse.ArgumentParser(description="从 .pptx 剥出空白底版")
    ap.add_argument("src", help="源 .pptx")
    ap.add_argument("-o", "--out", required=True, help="输出底版路径")
    args = ap.parse_args()

    if not os.path.exists(args.src):
        sys.exit(f"源文件不存在: {args.src}")

    n, layouts, masters = strip_slides(args.src, args.out)
    size_mb = os.path.getsize(args.out) / 1024 / 1024

    print(f"✓ {args.out}  ({size_mb:.2f} MB)")
    print(f"  已删除 {n} 张幻灯片，保留 {masters} 套母版 / {layouts} 个版式")

    if layouts <= 1:
        print("  ! 版式数量偏少，生成的课件可能缺少可用的版式骨架")
    if size_mb > 5:
        print("  ! 底版偏大（>5MB），通常是被删除幻灯片引用的媒体残留。"
              "可作为底版使用，但生成的每份课件都会带上这个体积")


if __name__ == "__main__":
    main()
