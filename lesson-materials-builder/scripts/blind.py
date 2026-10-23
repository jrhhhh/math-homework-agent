# -*- coding: utf-8 -*-
"""盲评装置 —— 判**脚本量不到**的那些维度。

为什么必须有它：`SKILL.md` 的核心主张是「按学生的**认知台阶**切页，而不是按教案
标题」。`references/deck-rules.md` 的 9 条规则里，规则 1、2、8 以及规则 4 的
「每台阶最后一页」判断，**代码零实现**。所以基准的盲区是被测物的镜像 ——
量得到的全是格式（标题字数、要点行数、口语词典），量不到的是它最想说的那件事。

试过给它造脚本代理，实测都被否掉：
  · 页标题 ↔ 教案环节名重合率：好 22% / 坏 20%，命中全是任何课件都有的尾页，
    换分词口径还会反号。
  · 纯文本语体规则：有技能 6 处 / 基线 7 处，打平，还把技能自己的阐述页误伤 6 次。
  · 告警计数：基线 58 条，但人写的比赛稿也有 34 条 —— 那量的是「是不是用这个技能
    做的」，不是「好不好」。
所以这一层不再造代理，交给模型判。**它贵，所以有闸门。**

三段式，和产出层一样（脚本起不了子代理，人在中间）：

    python scripts/bench.py blind-prep  --run <run 目录> --pair 名字=A.pptx:B.pptx ...
    #   ← 脱敏副本 + 导出每页 PNG + plan.json；然后由调用方起判官子代理
    python scripts/bench.py blind-grade --run <run 目录>
    #   ← 读 judgments.json，算位置一致性 / 好看对照 / 长度偏置 / 无法判断率

**脱敏是前置条件，不是收尾工作。** 实测：好/坏两份祖暅课件的封面都写着
`讲师：<真名>`，参赛稿的 core.xml 里还有两处元数据。**本文件自己不许写这个名字** ——
`lint_privacy.py` 扫的就是技能包，第一版草稿正是在这里栽的。喂给判官的副本一律生成在
run 目录（技能包之外），生成后**复扫一遍标识符，必须 0 命中** —— 这一条没过
就不导出、不起判官。
"""
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

# 判据全部取自 SKILL.md **自己声明的主张**，不是通用 PPT 质量。
# 通用质量（好看、整齐）由 CONTROL 那一题单独问，用来发现「其实在量好看」。
DIMENSIONS = [
    ("D1", "切页依据：这两份课件切页的依据是**学生的认知台阶**（每页对应一个思维"
           "动作：感受 / 比较 / 抽象 / 应用），还是**教案里的小节标题**？哪一份更"
           "接近前者？"),
    ("D2", "自检标准：学生扫一眼某一页，能不能说出「这页要我干什么」？哪一份更容易"
           "看出一页要干什么？"),
    ("D3", "驱动问题：课件的驱动问题是在**推动学生思考**（需要想一下才能答），还是"
           "**让学生复述**已经讲过的结论？哪一份更好？"),
    ("D4", "配图：图是在**支撑理解**（看得出这个图想说明什么），还是只是装饰 / 与"
           "这页要讲的事关系不大？哪一份的图更有用？"),
    ("D5", "要素完整性：教案里的环节（导入 / 探究 / 例题 / 小结 / 作业等）有没有被"
           "漏掉？哪一份更完整？"),
    ("CONTROL", "**只看好看程度**：哪一份视觉上更顺眼、更整齐？（这一题刻意不问"
                "认知台阶 —— 它的用处是当对照：如果它和 D1 选出同一个赢家，"
                "说明 D1 其实在量好看，D1 要记 NOT-RUN。）"),
]

VERDICT_RULES = {
    "position_inconsistent_max": 0.30,   # 换位后结论翻面超过这个比例 → 该维度作废
    "undecided_max": 0.50,               # 「无法判断」超过这个比例 → 该维度作废
}

# ---------------------------------------------------------------- 比赛评分表
#
# 两条轨，**分开报，不许混**：
#
#   A 轨「标准」—— 老师拿的这张比赛评分表。它是**外部**标准（不是从本技能的
#     脚本里长出来的），所以能当裁定用。但它里面**没有技能的核心主张**：
#     「内容」那条只说到「结构合理、逻辑顺畅、层次性和连贯性」—— 那是**结果
#     描述，不是判据**，一份按认知台阶切的和一份把教案大纲一节搬一页的，照它
#     读都可能算「结构合理」。而「驱动问题」（D3）整张表一个字没提。
#
#   B 轨「主张」—— D1 切页依据 / D3 驱动问题。这两条是 `SKILL.md` 自己要证明
#     的事，标准不覆盖，所以**比赛得分高不等于做到了这件事**。分开报才看得见差。
#
# 档位说明：原表的档是重叠的（很好 40—30 与好 35—25 在 30–35 上重合；技术/艺术/
# 创新三档在 10–15 上几乎全重叠），同一分数落进不同档 ⇒ 打出来的分不可比。
# 所以这里**只收数字、不收档位**。
RUBRIC_ITEMS = [
    ("内容", 40, "主体突出，内容完整；作品内容能够清晰的表达素材的精要。"
                 "结构合理，逻辑顺畅，幻灯片之间具有层次性和连贯性，过度恰当、顺畅。"
                 "内在素材符合主题：PPT 背景图片、音乐、视频等各种运用符合当选主题。"),
    ("技术运用", 20, "作品中对各种工具的运用如：文本、图片、表格、图表、动画、音频、视频等"
                     "表现工具的运用，作品中可使用超链接。整部作品的播放流畅，运行稳定，无故障。"),
    ("艺术视觉", 20, "整体界面美观布局合理，层次分明，视觉效果佳，表现力和感染力强。"
                     "字体设计恰当，文字清晰，风格引人入胜。"),
    ("创新", 20, "表意新颖，构思独特，设计巧妙，具有想象力和表现力，具有较大的创新，"
                 "原创成分高，具有鲜明的个性。"),
]


# ------------------------------------------------------------------ 脱敏

def desensitize(src, dst, identifiers):
    """把 pptx 当 zip 拆开，在所有 XML 里替换标识符，再打回去。

    改的是 `.xml` 与 `.rels` 的正文，**也包括 core.xml** —— 参赛稿的
    `dc:creator` / `cp:lastModifiedBy` 里各有一处真名，那是元数据不是版面，
    导出成图看不见，但文件本身带着它。
    """
    n = 0
    with zipfile.ZipFile(src) as zin:
        items = [(i, zin.read(i.filename)) for i in zin.infolist()]
    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for info, data in items:
            if info.filename.endswith((".xml", ".rels")):
                t = data.decode("utf-8", "replace")
                for ident in identifiers:
                    c = t.count(ident)
                    if c:
                        n += c
                        t = t.replace(ident, "教师")
                data = t.encode("utf-8")
            zout.writestr(info, data)
    return n


def rescan(path, identifiers):
    """复扫一份 pptx 里还有没有标识符。返回命中列表 —— 必须为空才算脱敏成功。"""
    hits = []
    with zipfile.ZipFile(path) as z:
        for nm in z.namelist():
            if not nm.endswith((".xml", ".rels")):
                continue
            t = z.read(nm).decode("utf-8", "replace")
            for ident in identifiers:
                if ident in t:
                    hits.append("%s:%s ×%d" % (os.path.basename(path), nm, t.count(ident)))
    return hits


# ------------------------------------------------------------------ 导图

PS_EXPORT = r'''
param([string]$Pptx, [string]$OutDir)
$ErrorActionPreference = "Stop"
if (-not (Test-Path $OutDir)) { New-Item -ItemType Directory -Path $OutDir | Out-Null }
$app = $null
foreach ($prog in @("PowerPoint.Application", "KWPP.Application")) {
  try { $app = New-Object -ComObject $prog; break } catch { }
}
if ($null -eq $app) { Write-Output "NO-COM"; exit 3 }
try {
  $pres = $app.Presentations.Open($Pptx, $true, $false, $false)
  foreach ($s in $pres.Slides) {
    $f = Join-Path $OutDir ("slide{0:D2}.png" -f $s.SlideIndex)
    $s.Export($f, "PNG", 1600, 900)
  }
  Write-Output ("EXPORTED " + $pres.Slides.Count)
  $pres.Close()
} finally {
  try { $app.Quit() } catch { }
}
'''


def export_slides(pptx, outdir):
    """把每一页导成 PNG。返回 (ok, 说明)。

    导的是**点击后**的样子（PowerPoint 与 WPS 皆如此，已实测）—— 对盲评无所谓，
    判官要看的是这页最终长什么样。
    """
    fd, ps = tempfile.mkstemp(suffix=".ps1", prefix="blind-export-")
    os.close(fd)
    with open(ps, "w", encoding="utf-8") as f:
        f.write(PS_EXPORT)
    try:
        p = subprocess.run(["powershell", "-NoProfile", "-File", ps,
                            "-Pptx", os.path.abspath(pptx),
                            "-OutDir", os.path.abspath(outdir)],
                           capture_output=True)
        out = (p.stdout + p.stderr).decode("utf-8", "replace").strip()
        if p.returncode != 0 or not out.startswith("EXPORTED"):
            return False, out[-300:] or "退出码 %d" % p.returncode
        return True, out
    except FileNotFoundError:
        return False, "找不到 powershell —— 导出不了，这一层的视觉维度只能记 NOT-RUN"
    finally:
        try:
            os.remove(ps)
        except OSError:
            pass


# ------------------------------------------------------------------ prep

def cmd_prep(a):
    run = a.run
    if not os.path.isdir(run):
        print("没有这个 run 目录：%s" % run)
        return 2
    idents = [x for x in (a.identifier or []) if x]
    if not idents:
        env = os.environ.get("BENCH_IDENTIFIERS", "")
        idents = [x.strip() for x in env.split(",") if x.strip()]
    if not idents:
        print("必须给标识符清单 —— 脱敏是这一层的前置条件，不是可选项。")
        print("  BENCH_IDENTIFIERS=真名,学号 python scripts/bench.py blind-prep ...")
        return 2

    pairs = []
    for spec in (a.pair or []):
        if "=" not in spec or ":" not in spec.split("=", 1)[1]:
            print("--pair 要写成 名字=A.pptx:B.pptx，收到的是：%s" % spec)
            return 2
        name, rest = spec.split("=", 1)
        pa, pb = rest.split(":", 1)
        pairs.append((name, os.path.abspath(pa), os.path.abspath(pb)))
    if not pairs:
        print("至少给一对：--pair 名字=A.pptx:B.pptx")
        return 2

    root = os.path.join(run, "blind")
    decks = os.path.join(root, "decks")
    slides = os.path.join(root, "slides")
    for d in (decks, slides):
        if os.path.isdir(d):
            shutil.rmtree(d)
        os.makedirs(d)

    # ---- ① 脱敏，并且**当场复扫**
    print("① 脱敏（副本只生成在 run 目录里，技能包不碰）")
    all_hits = []
    for name, pa, pb in pairs:
        for side, src in (("A", pa), ("B", pb)):
            if not os.path.isfile(src):
                print("   找不到 %s" % src)
                return 2
            dst = os.path.join(decks, "%s-%s.pptx" % (name, side))
            n = desensitize(src, dst, idents)
            hits = rescan(dst, idents)
            all_hits += hits
            print("   %-10s %s  替换 %d 处，复扫 %s"
                  % (name, side, n, "干净" if not hits else "**还有 %d 处！**" % len(hits)))
    if all_hits:
        print("\n脱敏没通过，**别往下走、更别起判官**：")
        for h in all_hits:
            print("   ", h)
        return 1
    print("   → 全部复扫 0 命中")

    # ---- ② 导出每页 PNG
    print("\n② 导出每页 PNG（判官看的是这个）")
    deck_files = []
    for name, pa, pb in pairs:
        for side, orig in (("A", pa), ("B", pb)):
            src = os.path.join(decks, "%s-%s.pptx" % (name, side))
            out = os.path.join(slides, "%s-%s" % (name, side))
            ok, msg = export_slides(src, out)
            n = len([f for f in os.listdir(out) if f.endswith(".png")]) if os.path.isdir(out) else 0
            print("   %-10s %s  %s（%d 张）" % (name, side, "好" if ok else "**失败**：" + msg, n))
            if ok:
                # `src_orig` 只记在 plan.json 里备查 —— 判官从目录名（`pair1-A`）
                # 看不出这份是谁做的。**别把这个字段喂给判官。**
                deck_files.append({"pair": name, "side": side,
                                   "dir": os.path.relpath(out, run).replace("\\", "/"),
                                   "slides": n, "src_orig": orig})

    # ---- ③ plan.json：判官要照它判
    exp = {}
    for spec in (getattr(a, "expect", None) or []):
        if "=" in spec:
            k, v = spec.split("=", 1)
            exp[k.strip()] = v.strip().upper()
    plan = {
        "pairs": [{"pair": n, "expect": exp.get(n), "A": pa, "B": pb}
                  for n, pa, pb in pairs],
        "dimensions": [{"id": i, "question": q} for i, q in DIMENSIONS],
        "rubric": [{"id": i, "max": m, "text": t} for i, m, t in RUBRIC_ITEMS],
        "orders": ["AB", "BA"],
        "deck_files": deck_files,
        "verdict_rules": VERDICT_RULES,
        "how_to_run": [
            "对每一对、每一个维度、每一种顺序各起一个判官子代理（判官 ≠ 生成者）。",
            "给它的输入只有：两个目录里的 PNG（按 orders 指定的先后），加上该维度的问题。",
            "**不要告诉它哪份是技能做的、哪份是基线的。**",
            "让它回一段 JSON：[{\"pair\",\"dim\",\"order\",\"winner\":\"A|B|无法判断\",",
            "  \"why\",\"made_by\":\"猜是谁做的\"}]，攒进 run/blind/judgments.json。",
        ],
    }
    with open(os.path.join(root, "plan.json"), "w", encoding="utf-8") as f:
        json.dump(plan, f, ensure_ascii=False, indent=1)
    print("\n③ plan.json 写好了：%s" % os.path.join(root, "plan.json"))
    print("   下一步：起判官子代理，把结果攒进 %s" % os.path.join(root, "judgments.json"))
    print("   然后：python scripts/bench.py blind-grade --run %s" % run)
    return 0


# ------------------------------------------------------------------ grade

def _tally(rows, dim):
    """一个维度上的读数。"""
    rs = [r for r in rows if r["dim"] == dim]
    orders = {}
    for r in rs:
        orders.setdefault(r["order"], []).append(r)
    undecided = sum(1 for r in rs if r.get("winner") in ("无法判断", "undecided", None))
    # 位置一致性：同一对在 AB / BA 两种顺序下选出的是**同一份课件**吗？
    # 注意比的是「哪一份」，不是「A 还是 B」—— 换了顺序，赢家该还是那一份。
    flip = same = 0
    for pair in {r["pair"] for r in rs}:
        ab = [r for r in orders.get("AB", []) if r["pair"] == pair]
        ba = [r for r in orders.get("BA", []) if r["pair"] == pair]
        if not ab or not ba:
            continue
        w_ab, w_ba = ab[0].get("winner"), ba[0].get("winner")
        if w_ab in ("无法判断", None) or w_ba in ("无法判断", None):
            continue
        # AB 下选 A = 第一份；BA 下选 B = 第一份。都换算成「第一份 / 第二份」之外
        # 的绝对身份：AB 的 A 与 BA 的 B 是同一份。
        same_ab = w_ab
        same_ba = {"A": "B", "B": "A"}.get(w_ba, w_ba)
        if same_ab == same_ba:
            same += 1
        else:
            flip += 1
    judged = same + flip
    return {
        "dim": dim, "n": len(rs), "undecided": undecided,
        "flip": flip, "same": same, "judged": judged,
        "flip_rate": (flip / judged) if judged else None,
        "undecided_rate": (undecided / len(rs)) if rs else None,
    }


def _load_plan(root):
    p = os.path.join(root, "plan.json")
    try:
        with open(p, encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return {}


def _human_scores(root):
    """读老师按比赛表打的分：`<run>/blind/人裁定.json`。

    ```json
    {"pair1": {"A": {"内容": 30, "技术运用": 15, "艺术视觉": 14, "创新": 12},
               "B": {...}}}
    ```
    缺文件、缺某一对、缺某一项，都只是那一格标不出来 —— **不报错、也不当通过**。
    """
    try:
        with open(os.path.join(root, "人裁定.json"), encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return {}


def decks_dir(run):
    return os.path.join(run, "blind", "decks")


def deck_chars(ddir):
    """数每份课件里所有文字加起来多少字。用来查判官有没有「奖励字多」。"""
    out = {}
    if not os.path.isdir(ddir):
        return out
    try:
        from pptx import Presentation
    except ImportError:
        return out
    for f in sorted(os.listdir(ddir)):
        if not f.endswith(".pptx"):
            continue
        try:
            prs = Presentation(os.path.join(ddir, f))
        except Exception:
            continue
        n = 0
        for s in prs.slides:
            for sh in s.shapes:
                if sh.has_text_frame:
                    n += len(re.sub(r"\s", "", sh.text_frame.text))
        out[f[:-5]] = n
    return out


def cmd_grade(a):
    run = a.run
    root = os.path.join(run, "blind")
    jp = os.path.join(root, "judgments.json")
    if not os.path.isfile(jp):
        print("还没有 judgments.json：%s" % jp)
        print("先 blind-prep，起判官子代理，把它们的结论攒成这个文件。")
        return 2
    with open(jp, encoding="utf-8") as f:
        rows = json.load(f)
    if isinstance(rows, dict):
        rows = rows.get("judgments", [])

    print("盲评判读 —— %d 条判断" % len(rows))
    print("=" * 64)
    verdicts = {}
    for dim, q in DIMENSIONS:
        t = _tally(rows, dim)
        if not t["n"]:
            verdicts[dim] = ("NOT-RUN", "这个维度一条判断都没有")
            continue
        if t["undecided_rate"] is not None and t["undecided_rate"] > VERDICT_RULES["undecided_max"]:
            verdicts[dim] = ("NOT-RUN",
                             "「无法判断」占 %.0f%%（>%.0f%%）—— 判官在这题上给不出信息，"
                             "**不许取平均当平局**" % (t["undecided_rate"] * 100,
                                                   VERDICT_RULES["undecided_max"] * 100))
            continue
        if t["flip_rate"] is not None and t["flip_rate"] > VERDICT_RULES["position_inconsistent_max"]:
            verdicts[dim] = ("NOT-RUN",
                             "位置不一致率 %.0f%%（>%.0f%%）—— 换一下左右就翻面，"
                             "位置偏置压过了信号" % (t["flip_rate"] * 100,
                                                 VERDICT_RULES["position_inconsistent_max"] * 100))
            continue
        verdicts[dim] = ("OK", "位置一致 %d/%d" % (t["same"], t["judged"]))

    for dim, q in DIMENSIONS:
        st, why = verdicts[dim]
        t = _tally(rows, dim)
        print("\n%s  %s" % (dim, "—— " + why if st == "NOT-RUN" else why))
        if st == "NOT-RUN":
            print("      %s" % q[:70])
        for pair in sorted({r["pair"] for r in rows if r["dim"] == dim}):
            ws = [r.get("winner") for r in rows if r["dim"] == dim and r["pair"] == pair]
            print("      %-10s → %s" % (pair, " ｜ ".join(str(w) for w in ws)))

    # ---- 好看对照：CONTROL 与 D1 选出同一个赢家 → D1 其实在量好看
    print("\n" + "=" * 64)
    print("对照：D1 是不是其实在量「好看」")
    bad_control = False
    for pair in sorted({r["pair"] for r in rows}):
        d1 = [r.get("winner") for r in rows if r["dim"] == "D1" and r["pair"] == pair]
        ct = [r.get("winner") for r in rows if r["dim"] == "CONTROL" and r["pair"] == pair]
        if d1 and ct and d1[0] == ct[0] and d1[0] not in ("无法判断", None):
            bad_control = True
            print("   ⚠ %s：D1 与「只判好看」选了同一份（%s）" % (pair, d1[0]))
    if bad_control:
        print("   → D1 记 NOT-RUN：它与「只判好看」分不开，不能当认知台阶的证据。")
        verdicts["D1"] = ("NOT-RUN", "与好看对照选了同一份")
    else:
        print("   ✓ D1 与好看对照选了不同的赢家（或有一方无法判断）—— D1 与好看可分离")

    # ---- 长度偏置：赢的那份是不是只是字多
    #
    # 字数由**脚本数**，不问判官 —— 判官数不准，而这一项要的正是准数。
    #
    # ⚠ `winner` 是「先给的那份」，**不是固定的哪一份课件**。AB 顺序下的 "A"
    # 和 BA 顺序下的 "A" 指的是两份不同的课件。所以这里必须按 order 换算成
    # 拷贝的绝对名（`pair1-A` / `pair1-B`）再去查字数 —— 早先的写法直接拿
    # `winner == "A"` 去对 `pair1-A` 的字数，把 BA 那一半全算反了，于是
    # 「12 条里赢家字更多 6 条」这个「没有偏置」的结论是假的。
    print("\n长度偏置：赢的那份是不是只是字多")
    chars = deck_chars(decks_dir(run))
    lb = []
    for r in rows:
        w = r.get("winner")
        if w not in ("A", "B"):
            continue
        first = "%s-%s" % (r["pair"], w)                  # 先给的那份
        a, b = "%s-A" % r["pair"], "%s-B" % r["pair"]
        if r.get("order") == "BA":
            first = {"%s-A" % r["pair"]: b, "%s-B" % r["pair"]: a}[first]
        other = b if first == a else a
        if chars.get(first) and chars.get(other):
            lb.append((chars[first], chars[other]))
    if lb:
        win_more = sum(1 for w, l in lb if w > l)
        avg_w = sum(w for w, _ in lb) / len(lb)
        avg_l = sum(l for _, l in lb) / len(lb)
        print("   %d 条里赢家字更多的有 %d 条（%.0f%%）；赢家均 %.0f 字 / 输家均 %.0f 字"
              % (len(lb), win_more, 100.0 * win_more / len(lb), avg_w, avg_l))
        pairs_n = len({r["pair"] for r in rows})
        if len(lb) and (win_more == len(lb) or win_more == 0):
            print("   → 一边倒。**但只有 %d 对课件** —— 字数在这 %d 对上和「哪份赢」"
                  "完全共线，拆不开「判官偏爱字多」和「字多的那份确实更好」。"
                  "要拆开得多给几对、且每对里两份字数互有高低。" % (pairs_n, pairs_n))
        elif win_more * 2 == len(lb):
            print("   → 对半，与字数无关。")
    else:
        print("   数不出字数，这一项 NOT-RUN（不是通过）")

    # ---- 识破指纹
    print("\n判官识破技能指纹了吗")
    mb = [r.get("made_by") for r in rows if r.get("made_by")]

    def _short(v):
        """判官的 made_by 是自由文本，有的写「A」有的写了一整段理由。
        只取起头那个判断（A / B / 两份都是 / 无法判断），别让一段话撑破版面。"""
        head = re.split(r"[（(（]|——|--", v, 1)[0].strip()
        return head[:18] or v[:18]

    if mb:
        # 和长度偏置一样，A/B 在这里也是「先给的那份」—— 换算成绝对身份，
        # 否则同一个 "A" 会把两份课件的票混在一起，看着像有倾向其实是噪声。
        tally = {}
        for r in rows:
            v = r.get("made_by")
            if not v:
                continue
            h = _short(v)
            if h in ("A", "B") and r.get("order") == "BA":
                h = {"A": "B", "B": "A"}[h]
            if h in ("A", "B"):
                h = "指认 %s（%s）" % (h, r["pair"])
            tally[h] = tally.get(h, 0) + 1
        for v, c in sorted(tally.items(), key=lambda kv: -kv[1]):
            print("   %-26s ×%d" % (v, c))
        print("   （判官原文见 judgments.json 的 made_by 字段）")
        picked = {k: v for k, v in tally.items() if k.startswith("指认")}
        if picked and max(picked.values()) * 2 >= len(mb):
            print("   → 过半判官指认同一份 —— 盲没做到，这一对的结论要打折")
        else:
            print("   → 指认分散、没有过半的共识 —— 判官认不出哪份是工具做的，盲站得住")
    else:
        print("   判官没回「猜是谁做的」，这一项 NOT-RUN")

    print("\n" + "=" * 64)
    ok = [d for d, (s, _) in verdicts.items() if s == "OK"]
    nr = [d for d, (s, _) in verdicts.items() if s == "NOT-RUN"]
    print("可用维度 %d：%s" % (len(ok), "、".join(ok) or "无"))
    print("作废维度 %d：%s" % (len(nr), "、".join(nr) or "无"))

    # ---- 闸门。**判的是「判官有没有选出那个本该赢的」，不是「它有没有选出
    # 一个赢家」。** 两者天差地别：判官可能非常干脆地、两种顺序都一致地选出
    # **反的那一份** —— 那说明的不是判官不行，是**参照本身标错了**。
    # 早先这里的结语是写死的一段话，它只会念「判官分不开已知的好与已知的坏」，
    # 而实测那次 12 票全投给了标着「坏」的那份。结论句子和事实相反 ——
    # 正是这个项目一路在防的那种事，只是这次发生在结语里。
    # ---- A 轨：比赛评分表。**和老师打的分数对。**
    #
    # 这一轨才是「判官对不对」，上面那些只说明「判官稳不稳」。稳和对是两件事：
    # 一个两种顺序都一致地判错的判官，位置一致性是 100%。
    print("\n" + "=" * 64)
    print("A 轨「标准」—— 比赛评分表")
    rub = [r for r in rows if r.get("dim") == "RUBRIC" and r.get("scores")]
    human = _human_scores(root)
    if not rub:
        print("   判官没打这张表 —— 这一轨 NOT-RUN（不是通过）")
    else:
        for pair in sorted({r["pair"] for r in rub}):
            rows_p = [r for r in rub if r["pair"] == pair]
            print("\n   %s" % pair)
            h = human.get(pair) or {}
            print("      %-10s %7s %7s %5s %5s" %
                  ("项", "判-A", "判-B", "人A", "人B"))
            signs, agree = [], 0
            for item, mx, _ in RUBRIC_ITEMS:
                # ⚠ 判官回的 A/B 是**先给的那份**，两种顺序下指的不是同一份
                # 课件。必须按 order 换算成拷贝的绝对名再平均 —— 直接把
                # 两种顺序的 `scores["A"]` 混在一起平均，得到的是两份不同
                # 课件的平均，那个数没有任何意义（而且看起来很像一个数）。
                # 同类错在这个文件里犯过三次（长度偏置、made_by、这里）。
                def avg(deck):
                    vs = []
                    for r in rows_p:
                        side = {"AB": {"A": "A", "B": "B"},
                                "BA": {"A": "B", "B": "A"}}.get(
                                    r.get("order"), {}).get(deck)
                        if side is None:
                            continue
                        v = r["scores"].get(side, {}).get(item)
                        if isinstance(v, (int, float)):
                            vs.append(v)
                    return sum(vs) / len(vs) if vs else None
                ja, jb = avg("A"), avg("B")
                ha = (h.get("A") or {}).get(item)
                hb = (h.get("B") or {}).get(item)
                def fmt(v):
                    return "—" if v is None else ("%g" % round(v, 1))
                note = ""
                if None not in (ja, jb) and isinstance(ha, (int, float)) \
                        and isinstance(hb, (int, float)) and ha != hb and ja != jb:
                    js = 1 if ja > jb else -1
                    hs = 1 if ha > hb else -1
                    signs.append(js == hs)
                    if js == hs:
                        agree += 1
                        note = "一致"
                    else:
                        note = "**反了**"
                print("      %-10s %5s %5s %5s %5s   %s   （满分 %d）"
                      % (item, fmt(ja), fmt(jb), fmt(ha), fmt(hb), note, mx))
            if signs:
                print("      → 这一对上，判官和人同向 %d/%d 项" % (agree, len(signs)))
            else:
                print("      → 人对这一对还没打分（或两项打平），这一对**标定不了**")

    print("\n" + "=" * 64)
    print("**闸门**")
    pairs_exp = {p["pair"]: p.get("expect") for p in _load_plan(root).get("pairs", [])
                 if p.get("expect")}
    if not pairs_exp:
        print("   没给参照（blind-prep --expect 名字=A）—— 闸门**判不了**，记 NOT-RUN。")
        print("   判官选出了谁不等于选对了谁；没有「本该赢的是哪份」这个信息，")
        print("   就无法区分「判官不行」和「参照标错了」。")
    else:
        for pair, want in sorted(pairs_exp.items()):
            print("   %s：参照认为 %s 更好" % (pair, want))
            for dim, _ in DIMENSIONS:
                votes = []
                for r in rows:
                    if r["dim"] != dim or r["pair"] != pair:
                        continue
                    w = r.get("winner")
                    if w not in ("A", "B"):
                        continue
                    # 换算成拷贝的绝对身份（A/B 是「先给的那份」，随 order 变）
                    if r.get("order") == "BA":
                        w = {"A": "B", "B": "A"}[w]
                    votes.append(w)
                if not votes:
                    continue
                agree = sum(1 for v in votes if v == want)
                mark = "✓" if agree * 2 > len(votes) else "✗"
                print("      %s %-8s %d/%d 票投给 %s"
                      % (mark, dim, agree, len(votes), want))
        print("\n   ✗ 占多/全部 → **判官选的和你标的参照相反**。这不等于判官不行：")
        print("   先去核参照本身 —— 它是按什么定的？如果参照的「好/坏」是从脚本")
        print("   告警数来的，那就还是循环论证，闸门不成立，记 NOT-RUN 并停手。")

    out = os.path.join(root, "blind-verdict.json")
    with open(out, "w", encoding="utf-8") as f:
        json.dump({"verdicts": {d: {"status": s, "why": w} for d, (s, w) in verdicts.items()},
                   "tallies": {d: _tally(rows, d) for d, _ in DIMENSIONS},
                   # A 轨的分也是结果的一部分 —— 只写进 stdout 就等于下次重跑
                   # 再也拿不回来，而报告要引它。
                   "rubric_scores": [r for r in rows
                                     if r.get("dim") == "RUBRIC" and r.get("scores")],
                   "human_scores": _human_scores(root)},
                  f, ensure_ascii=False, indent=1)
    print("\n→ %s" % out)
    return 0
