#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""变异矩阵：把基准课件逐条改坏，每条必须触发一个**指定**的错误。

    python evals/fixtures/mutate_pptx.py 基准.pptx 输出目录 [--check]

## 为什么是「改坏 pptx」而不是「手写 .md」

`verify_pptx` 查的那些几何毛病 —— 图片被拉伸、越出画布、间距不均、行距写成
百分比 —— **渲染器压根产不出来**。`build_pptx` 永远给出正确几何。所以拿 `.md`
当夹具是无效的：夹具本身就会渲染成一份正确课件，测不到任何东西。

要测「检查器认不认得这个毛病」，就必须**在渲染器下游注入缺陷**：先拿一份
0 ERROR 的基准课件，再针对性地破坏一处属性。

## 为什么每条都配一句「必须出现」

这是整套设计的骨干。一个**见谁都报**的检查等于没有检查，一个**静默失效**的
检查更糟 —— 报告上都是绿的。所以两个方向都钉死：

  · 基准课件必须 0 ERROR（否则下面每条断言都失去意义）
  · 每条变异必须报出**它自己那一类**错误（用子串点名，不只看计数）

哪条检查哪天被关掉、或者被调宽了容差，红的是它自己的那一格。
"""

import os
import shutil
import sys

# Windows 上 stdout 默认 GBK，打一个 ✓ 就 UnicodeEncodeError 崩掉。
# verify_pptx.py 当初正是栽在这上面 —— 一份**完全合格**的课件被判成失败。
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

from pptx import Presentation
from pptx.util import Emu, Inches, Pt

IN = 914400.0

# 名字 → (说明, 变异函数, 必须出现的错误子串)
# 变异函数就地改 Presentation。
MUTATIONS = {}


def mutation(name, expect, doc, xfail=False):
    """xfail=True：当前**预期不报**。不报算通过；哪天报了会打 XPASS 提醒摘掉。

    与那 3 条词典漏洞的处理一致 —— 让漏检**可见**，而不是假装它不存在。
    """
    def deco(fn):
        MUTATIONS[name] = (doc, fn, expect, xfail)
        return fn
    return deco


# ---------- 小工具 ----------

def _pics(slide):
    return [s for s in slide.shapes if s.shape_type == 13]


def _textboxes(slide):
    return [s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]


def _body(slide):
    """正文块：够高、段落最多的那个文本框。

    别用 `_textboxes(slide)[0]` —— 那是**标题**。标题框高 0.79in，够不着
    verify 里「要点条数」「行距」那几条检查的 1.0in 门槛，变异会静默失效，
    矩阵报「一条都没报」，看着像检查是哑的，其实是变异打偏了。
    """
    cands = [s for s in _textboxes(slide) if s.height >= 1.0 * IN]
    return max(cands, key=lambda s: len(s.text_frame.paragraphs)) if cands else None


def _text_left_of(slide, pic):
    """图左边那一栏文字块 —— 用来把图精确推到文字上。"""
    best = None
    for t in _textboxes(slide):
        right = t.left + t.width
        if right > pic.left + 0.05 * IN:          # 不在图的左边
            continue
        if t.top >= pic.top + pic.height or t.top + t.height <= pic.top:
            continue                               # 竖直不重叠
        if best is None or right > best.left + best.width:
            best = t
    return best


def _runs(slide):
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if r.text.strip():
                    yield sh, p, r


# ---------- 画布 ----------

@mutation("canvas-43", "画布", "把画布改成 4:3 —— 投影会整体变形")
def _m(prs):
    prs.slide_width = Emu(9144000)
    prs.slide_height = Emu(6858000)


# ---------- 位置与比例 ----------

@mutation("offscreen", "越出画布", "把一张图推出画布右缘")
def _m(prs):
    p = _pics(prs.slides[1])[0]
    p.left = Emu(prs.slide_width - int(0.4 * IN))


@mutation("stretch", "图片被拉伸", "只改宽不改高 —— 图会横向拉长")
def _m(prs):
    p = _pics(prs.slides[1])[0]
    p.width = Emu(int(p.width * 1.35))


@mutation("offcenter", "未居中", "整排图左移，排内间距不变")
def _m(prs):
    d = int(0.35 * IN)
    for p in _pics(prs.slides[1]):
        p.left = Emu(p.left - d)


@mutation("gap-uneven", "间距不均", "排内两图中移：外缘不动、居中不破，只改间距")
def _m(prs):
    pics = _pics(prs.slides[1])
    d = int(0.15 * IN)
    pics[1].left = Emu(pics[1].left + d)
    pics[2].left = Emu(pics[2].left + d)


@mutation("gap-negative", "图文间隙为负", "把右栏的图轻推一下，压到左栏文字列上")
def _m(prs):
    slide = prs.slides[3]
    p = _pics(slide)[0]
    t = _text_left_of(slide, p)
    # 分寸很要紧。推过头（比如 2in）图就跑到了文字块左缘之外，verify 的配对
    # 条件「文字右缘 ≤ 图左缘 + 0.05in」不再成立 → 那条检查直接 continue，
    # 只剩一条笼统的「图片与文字可能重叠」WARN。要让「间隙为负」这条 ERROR
    # 真的报出来，得把图停在文字右缘**之内**、配对仍成立的窄带里。
    p.left = Emu(t.left + t.width - int(0.02 * IN))


# ---------- 字体 ----------

@mutation("font-break", "字体不在契约内", "a:latin 换成契约外的 Arial")
def _m(prs):
    for sh, para, r in _runs(prs.slides[2]):
        r.font.name = "Arial"
        return


@mutation("times", "出现 Times New Roman", "中文课件的经典事故字体")
def _m(prs):
    for sh, para, r in _runs(prs.slides[2]):
        r.font.name = "Times New Roman"
        return


# ---------- 行距 ----------

@mutation("spcpct", "行距写成百分比", "行距用 spcPct —— 字号一变行距就跟着变")
def _m(prs):
    from pptx.oxml.ns import qn
    for sh in _textboxes(prs.slides[2]):
        if sh.height < 1.0 * IN:
            continue
        for para in sh.text_frame.paragraphs:
            if not "".join(r.text for r in para.runs).strip():
                continue
            pPr = para._p.get_or_add_pPr()
            old = pPr.find(qn("a:lnSpc"))
            if old is not None:
                pPr.remove(old)
            lnSpc = pPr.makeelement(qn("a:lnSpc"), {})
            lnSpc.append(pPr.makeelement(qn("a:spcPct"), {"val": "150000"}))
            pPr.insert(0, lnSpc)
        return


# ---------- 内容 ----------

def _dup(prs, rename):
    import copy
    from pptx.oxml.ns import qn
    sh = _body(prs.slides[2])
    dup = copy.deepcopy(sh._element)
    if rename:
        cn = dup.find(".//" + qn("p:cNvPr"))
        if cn is not None:
            cn.set("name", "重复块")
    sh._element.getparent().append(dup)


@mutation("dup-text", "同页重复文本", "文本框被复制一份（名字跟着变，像 Ctrl+D 那样）")
def _m(prs):
    _dup(prs, rename=True)


# 原先是 xfail：判重靠 seen[norm].add(sh.name)，同名同文的两个块被并成一个
# 名字就不报 —— 而「渲染器把同一块吐了两次」恰恰必然同名，正是它最该抓的。
# 判据已改成按 id(sh) 计数，这条现在真的会报，xfail 摘掉。
@mutation("dup-text-samename", "同页重复文本", "同一块吐两次，名字也相同")
def _m(prs):
    _dup(prs, rename=False)


@mutation("footer-leak", "内页出现课题/教材出处", "内页混进教材出处")
def _m(prs):
    slide = prs.slides[2]
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(6.4), Inches(4.0), Inches(0.5))
    tb.text_frame.text = "人教A版必修一 4.2.2"


@mutation("body-rides-up", "上沿高于正文区顶", "正文块整个上移，骑到标题带上")
def _m(prs):
    """钉住「正文块上沿高于正文区顶」这条检查**没有被修哑**。

    这条变异是跟着一次误报修复一起加的，缺了它那次修复就不算数：
    verify 原来用「字号 > 40pt」认标题，把 30pt 的节标题误判成正文块，于是
    任何标题框 ≥1.0in 高的课件都会吃到假告警（实测：基线那份 10 页报了 9 页，
    而它的正文规规矩矩从 1.90in 起排，一个毛病都没有）。
    修法是**跳过每页最高的那个文字块**。

    可「跳过最高的块」离「这条检查整个不跑了」只差一步 —— 所以这里把正文块
    顶到 0.90in（正文区顶 1.42、容差 0.15，越过门槛），**同时让标题仍在它上面**
    （标题 0.39in），证明检查只放过了标题，没有放过正文。
    """
    # 用 P3 —— P2 是流程页（四张图 + 说明行），没有 ≥1.0in 的正文块，
    # `_body()` 会返回 None，变异打空。
    sh = _body(prs.slides[2])
    sh.top = Emu(int(0.90 * IN))


@mutation("bullets-7", "要点 7 行", "一级要点超过规则 7 的上限 6 行")
def _m(prs):
    sh = _body(prs.slides[2])
    tf = sh.text_frame
    for i in range(4, 8):
        p = tf.add_paragraph()
        p.text = "第 %d 条要点" % i
        for r in p.runs:
            r.font.size = Pt(18)


# ---------- 驱动器 ----------

def main(base, outdir, do_check=False):
    os.makedirs(outdir, exist_ok=True)
    manifest = []

    # 先确认基准是干净的 —— 否则下面每一条断言都失去意义
    if do_check:
        sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "..", "scripts"))
        from verify_pptx import check
        prs, errs, warns, info = check(base)
        if errs:
            print("基准课件本身就有 %d 条 ERROR，先修基准：" % len(errs))
            for e in errs:
                print("   ", e)
            return 1
        print("基准课件：0 ERROR ✓")

    for name in sorted(MUTATIONS):
        doc, fn, expect, xfail = MUTATIONS[name]
        out = os.path.join(outdir, "mut-%s.pptx" % name)
        shutil.copyfile(base, out)
        prs = Presentation(out)
        fn(prs)
        prs.save(out)
        manifest.append({"file": out, "name": name, "expect": expect,
                         "doc": doc, "xfail": xfail})

    if not do_check:
        for m in manifest:
            print("  %-19s → 应报「%s」%s"
                  % (m["name"], m["expect"], "（xfail）" if m["xfail"] else ""))
        return 0

    from verify_pptx import check

    def msg(x):
        """check() 返回的是 (页号, 消息) 元组，不是裸字符串。"""
        return x[1] if isinstance(x, tuple) else str(x)

    print()
    bad = xpass = 0
    for m in manifest:
        prs, errs, warns, info = check(
            m["file"], os.path.join(os.path.dirname(__file__), "figures"))
        # 有些变异针对的是 WARN 级检查（比如要点条数超限），不全是 ERROR
        hit = [e for e in (list(errs) + list(warns)) if m["expect"] in msg(e)]
        if hit and not m["xfail"]:
            print("  ✓ %-19s %s" % (m["name"], msg(hit[0])[:74]))
        elif hit and m["xfail"]:
            xpass += 1
            print("  ◆ %-19s XPASS —— 已知漏洞似乎修好了，可以把 xfail 摘掉" % m["name"])
        elif m["xfail"]:
            print("  ◦ %-19s xfail（仍如预期不报：%s）" % (m["name"], m["doc"]))
        else:
            bad += 1
            print("  ✗ %-19s 期望「%s」，实际 ERROR %d / WARN %d："
                  % (m["name"], m["expect"], len(errs), len(warns)))
            for e in list(errs)[:3] + list(warns)[:2]:
                print("        %s" % msg(e)[:86])
            if not errs and not warns:
                print("        （一条都没报 —— 这条检查是哑的）")
    print()
    n_ok = len(manifest) - bad
    print("变异矩阵：%d/%d 命中%s" % (n_ok, len(manifest),
                                    "，另有 %d 条 XPASS" % xpass if xpass else ""))
    return 1 if bad else 0


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if len(args) < 2:
        print(__doc__)
        sys.exit(2)
    sys.exit(main(args[0], args[1], "--check" in sys.argv))
