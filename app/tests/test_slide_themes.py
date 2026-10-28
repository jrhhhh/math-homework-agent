"""Native template application, editability and compatibility checks."""
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[2]
BUILDER = ROOT / 'lesson-materials-builder'


class NativeSlideThemeTest(unittest.TestCase):
    def build(self, source, out, *options):
        result = subprocess.run([sys.executable, str(BUILDER / 'scripts/build_pptx.py'),
                                 str(source), '-o', str(out), *map(str, options)],
                                capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        return Presentation(out)

    def test_themes_have_real_layouts_editable_text_and_notes(self):
        with tempfile.TemporaryDirectory() as folder:
            for theme, prefix in [('scholar-red', 'ScholarRed'), ('ink-teal', 'InkTeal')]:
                out = Path(folder) / (theme + '.pptx')
                prs = self.build(BUILDER / 'assets/themes/showcase.md', out, '--theme', theme)
                self.assertEqual(len(prs.slides), 8)
                self.assertEqual(len([l for l in prs.slide_layouts if l.name.startswith(prefix + '.')]), 8)
                self.assertTrue(all(s.slide_layout.name.startswith(prefix + '.') for s in prs.slides))
                self.assertTrue(all(s.notes_slide.notes_text_frame.text for s in prs.slides))
                self.assertFalse(any(s.shape_type == 13 for s in prs.slides[0].shapes))
                # Background geometry lives in p:sldLayout, and content remains native text.
                self.assertGreater(len(prs.slides[0].slide_layout.shapes), 3)
                texts = [sh.text for slide in prs.slides for sh in slide.shapes if sh.has_text_frame]
                self.assertTrue(any('sin(2x)' in t for t in texts))
                process = prs.slides[6]
                self.assertEqual(sum('统一角' in sh.text for sh in process.shapes if sh.has_text_frame), 1)
                check = subprocess.run([sys.executable, str(BUILDER / 'scripts/verify_pptx.py'), str(out)], capture_output=True, text=True)
                self.assertEqual(check.returncode, 0, check.stdout + check.stderr)

    def test_native_template_file_carries_its_theme(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / 'source.md'
            source.write_text('---\ntitle: 测试\n---\n# [封面] 测试\n', encoding='utf-8')
            prs = self.build(source, Path(folder) / 'result.pptx', '--template', BUILDER / 'assets/themes/ink-teal.pptx')
            self.assertEqual(prs.slides[0].slide_layout.name, 'InkTeal.Cover')

    def test_classic_override_preserves_legacy_path(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / 'source.md'
            source.write_text('---\ntitle: 测试\ntheme: scholar-red\n---\n# [封面] 测试\n', encoding='utf-8')
            prs = self.build(source, Path(folder) / 'classic.pptx', '--theme', 'classic')
            self.assertFalse(any(l.name.startswith('ScholarRed.') for l in prs.slide_layouts))

    def test_custom_template_is_not_restyled_without_explicit_theme(self):
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            template = Presentation()
            template.slide_width = Inches(13 + 1 / 3)
            template.slide_height = Inches(7.5)
            template.slide_layouts[0].name = 'OWN_STYLE'
            template.save(folder / 'custom.pptx')
            (folder / 'source.md').write_text('---\ntitle: 测试\n---\n# [封面] 测试\n', encoding='utf-8')
            prs = self.build(folder / 'source.md', folder / 'result.pptx', '--template', folder / 'custom.pptx')
            self.assertIn('OWN_STYLE', [l.name for l in prs.slide_layouts])
            self.assertFalse(any(l.name.startswith('ScholarRed.') for l in prs.slide_layouts))


if __name__ == '__main__':
    unittest.main()
