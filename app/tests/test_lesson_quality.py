"""Coverage for instructional guidance and the unit-circle rendering path."""
import importlib.util
import json
from pathlib import Path
import tempfile
import unittest

from app.lesson_studio import LessonStudio
from app.server import Store

ROOT = Path(__file__).resolve().parents[2]


class LessonQualityTest(unittest.TestCase):
    def test_guidance_reaches_both_model_stages_and_unit_circle_is_accepted(self):
        prompts = []
        responses = iter([
            {"outline_markdown": "# [封面] 三角方程", "missing_information": [], "clarifying_questions": []},
            {"slides_markdown": "# [封面] 三角方程", "figures_spec": {"figures": [
                {"name": "circle", "recipe": "unit_circle", "angles_degrees": [60, 300]}]}},
        ])
        def generator(prompt):
            prompts.append(prompt)
            return {"raw": json.dumps(next(responses), ensure_ascii=False)}
        with tempfile.TemporaryDirectory() as tmp:
            studio = LessonStudio(Store(Path(tmp) / "test.db"), generator)
            lesson = studio.create_outline({"topic": "三角方程", "grade": "高一", "textbook": "人教A版",
                                            "duration": "20分钟", "draft_text": "解 sin x=0。"})
            lesson = studio.generate_source(lesson["id"], {"outline": lesson["outline"]})
            self.assertEqual(lesson["figures_spec"]["figures"][0]["recipe"], "unit_circle")
            self.assertTrue(all("数学课件：内容和视觉审查" in p for p in prompts))
            self.assertIn("unit_circle", prompts[1])
            self.assertTrue(all("可授课的数学课件设计" in p for p in prompts))
            self.assertIn("concept-lesson.md", prompts[1])
            self.assertIn("!note:", prompts[1])

    def test_two_columns_with_unequal_text_and_circles_pass_geometry(self):
        responses = iter([
            {"outline_markdown": "# [封面] 三角方程", "missing_information": [], "clarifying_questions": []},
            {"slides_markdown": "---\ntitle: 三角方程\nversion: daily\n---\n# [封面] 三角方程\n# [两栏] 区间筛选\n!note: 教师专用：先问除式为零时怎么办。\n- 取值范围 0≤x<2π\n## 左\n- 第一支的条件\n- sin x=0，解得 x=0 或 π。\n!img: figures/a.png\n## 右\n- cos x=1/2\n!img: figures/b.png\n", "figures_spec": {"figures": [
                {"name": "a", "recipe": "unit_circle", "angles_degrees": [0, 180]},
                {"name": "b", "recipe": "unit_circle", "angles_degrees": [60, 300]}]}},
        ])
        with tempfile.TemporaryDirectory() as tmp:
            studio = LessonStudio(Store(Path(tmp) / "test.db"), lambda _: {"raw": json.dumps(next(responses), ensure_ascii=False)})
            lesson = studio.create_outline({"topic": "三角方程", "grade": "高一", "textbook": "人教A版", "duration": "20分钟", "draft_text": "解三角方程。"})
            lesson = studio.generate_source(lesson["id"], {"outline": lesson["outline"]})
            result = studio.render(lesson["id"], {"slides_markdown": lesson["slides_markdown"]})
            self.assertEqual(result["status"], "ready", result["result"])
            from pptx import Presentation
            deck = Presentation(studio.file_path(lesson["id"], "deck"))
            text = " ".join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
            self.assertIn("取值范围", text)
            self.assertNotIn("教师专用", text)
            self.assertIn("教师专用：先问除式为零时怎么办。", deck.slides[1].notes_slide.notes_text_frame.text)

    def test_unit_circle_output_and_invalid_angles(self):
        spec = importlib.util.spec_from_file_location("lesson_figures", ROOT / "lesson-materials-builder/scripts/make_figures.py")
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        with tempfile.TemporaryDirectory() as tmp:
            target = str(Path(tmp) / "circle.png")
            module.recipe_unit_circle({"angles_degrees": [60, 300]}, target)
            self.assertGreater(Path(target).stat().st_size, 1000)
            for angles in ([360], [-1], [1, 1], [float('nan')], list(range(9)), [True]):
                with self.assertRaises(ValueError):
                    module.recipe_unit_circle({"angles_degrees": angles}, target)


if __name__ == "__main__":
    unittest.main()
