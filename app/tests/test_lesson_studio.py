import base64
import json
from pathlib import Path
import tempfile
import unittest

from app.lesson_studio import LessonStudio, _extract_upload
from app.server import Store


class LessonStudioTest(unittest.TestCase):
    def test_text_upload_is_decoded_and_unsupported_file_is_rejected(self):
        text = base64.b64encode("第一行\n第二行".encode()).decode()
        self.assertEqual(_extract_upload("lesson.txt", text), "第一行\n第二行")
        with self.assertRaisesRegex(ValueError, "只支持"):
            _extract_upload("lesson.exe", text)

    def test_docx_upload_extracts_paragraphs(self):
        from docx import Document
        from io import BytesIO

        document = Document()
        document.add_paragraph("教学目标")
        document.add_paragraph("学生观察函数图象")
        buffer = BytesIO()
        document.save(buffer)
        encoded = base64.b64encode(buffer.getvalue()).decode()
        extracted = _extract_upload("lesson.docx", encoded)
        self.assertEqual(extracted, "教学目标\n学生观察函数图象")

    def test_outline_source_and_real_builder_pipeline(self):
        with tempfile.TemporaryDirectory() as folder:
            store = Store(Path(folder) / "workbench.db")
            responses = iter([
                {"outline_markdown": "# [封面] 测试课\n# [标题] 观察函数\n- 学生比较两种图象", "missing_information": [], "clarifying_questions": []},
                {"slides_markdown": "---\ntitle: 测试课\nversion: daily\n---\n# [封面] 测试课\n教材：人教A版\n# [标题] 观察函数\n- 学生比较两种图象\n<!-- 允许无图 -->", "figures_spec": {"figures": []}},
            ])

            def generator(_prompt):
                return {"raw": json.dumps(next(responses), ensure_ascii=False)}

            studio = LessonStudio(store, generator=generator)
            lesson = studio.create_outline({
                "topic": "函数图象观察", "grade": "高一", "textbook": "人教A版",
                "duration": "10分钟", "draft_text": "学生观察函数图象并归纳性质。",
            })
            self.assertEqual(lesson["status"], "outline_ready")
            lesson = studio.generate_source(lesson["id"], {"outline": lesson["outline"], "clarifications": ""})
            self.assertEqual(lesson["status"], "source_ready")
            result = studio.render(lesson["id"], {"slides_markdown": lesson["slides_markdown"]})
            self.assertEqual(result["status"], "ready", result["result"])
            deck = studio.file_path(lesson["id"], "deck")
            source = studio.file_path(lesson["id"], "source")
            self.assertTrue(deck and deck.stat().st_size > 1000)
            self.assertTrue(source and "测试课" in source.read_text(encoding="utf-8"))
            self.assertGreaterEqual(result["result"]["slides"], 2)

    def test_incomplete_outline_requires_teacher_clarification(self):
        with tempfile.TemporaryDirectory() as folder:
            store = Store(Path(folder) / "workbench.db")
            studio = LessonStudio(store, generator=lambda _prompt: {"raw": json.dumps({
                "outline_markdown": "# [封面] 测试", "missing_information": ["请补充例题来源"],
                "clarifying_questions": ["例题来自哪一节？"],
            }, ensure_ascii=False)})
            lesson = studio.create_outline({"topic": "测试", "grade": "高一", "textbook": "人教A版",
                                            "duration": "45分钟", "draft_text": "教学目标：理解函数。"})
            with self.assertRaisesRegex(ValueError, "补充大纲"):
                studio.generate_source(lesson["id"], {"outline": lesson["outline"], "clarifications": ""})


if __name__ == "__main__":
    unittest.main()
