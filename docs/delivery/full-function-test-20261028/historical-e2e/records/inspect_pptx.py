"""Inspect and render the actual generated PowerPoint using local tools."""
import hashlib,json,shutil,subprocess
from pathlib import Path
from pptx import Presentation
root=Path(__file__).resolve().parents[1]
decks=list(root.glob('lessons/*/课件.pptx'))
if len(decks)!=1:raise SystemExit(f'Expected one deck, found {len(decks)}')
deck=decks[0];slides=[]
for i,slide in enumerate(Presentation(deck).slides,1):
 text='\n'.join(shape.text for shape in slide.shapes if getattr(shape,'has_text_frame',False) and shape.text)
 slides.append({'page':i,'text':text})
result={'deck':str(deck.relative_to(root)),'sha256':hashlib.sha256(deck.read_bytes()).hexdigest(),'slides':slides,
        'scope':'python-pptx 提取文本；LibreOffice 静态转换截图不验证 PowerPoint 原生动画和播放效果。'}
preview=root/'课件静态预览';preview.mkdir(exist_ok=True)
if shutil.which('soffice'):
 result['conversion']={}
 command=['soffice','-env:UserInstallation='+ (root/'lo-profile').as_uri(), '--headless','--convert-to','pdf','--outdir',str(preview),str(deck)]
 proc=subprocess.run(command,capture_output=True,text=True,timeout=120)
 result['conversion']={'exit_code':proc.returncode,'stdout':proc.stdout,'stderr':proc.stderr}
 pdf=preview/'课件.pdf'
 if proc.returncode==0 and pdf.is_file() and shutil.which('pdftoppm'):
  image_cmd=['pdftoppm','-png','-r','90',str(pdf),str(preview/'幻灯片')]
  p=subprocess.run(image_cmd,capture_output=True,text=True,timeout=120)
  result['preview']={'exit_code':p.returncode,'images':[str(f.relative_to(root)) for f in sorted(preview.glob('幻灯片-*.png'))],'stderr':p.stderr[-2000:]}
Path(__file__).with_name('pptx-inspection.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'slides':len(slides),'conversion':result.get('conversion'),'preview':result.get('preview')},ensure_ascii=False))
